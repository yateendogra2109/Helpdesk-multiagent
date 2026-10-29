"""Build data/csv/courses.csv and timetable.csv from the department course-list workbook.

Sources:
  data/pdfs/Finalized courses list .xlsx   one sheet per department, layouts differ
  CURRENT_COURSES (below)                  transcribed from data/ss/current_courses.png
                                           (a student's registered courses; the workbook has no CSE sheet)
  data/csv/slot_system.csv                 slot letter -> weekly meeting times
  data/csv/instructors.csv                 faculty directory, used to link instructors to ids

Usage: python scripts/build_csvs.py
"""
import csv
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "data" / "pdfs" / "Finalized courses list .xlsx"
CSV_DIR = ROOT / "data" / "csv"
SEMESTER = "2026-I"

SHEET_DEPT = {
    "CCEI(SAIDE)": "School of Artificial Intelligence & Data Science",
    "SME": "School of Management & Entrepreneurship",
    "Economics": "Economics",
    "Math": "Mathematics",
    "CIE": "Civil and Infrastructure Engineering",
    "Energy & Automotive": "Energy and Automotive Engineering",
    "BSBE": "Bioscience & Bioengineering",
    "Chemistry": "Chemistry",
    "CET": "Centre for Education Technology",
    "ME": "Mechanical Engineering",
    "MT": "Materials Engineering",
    "SoLA": "School of Liberal Arts (Humanities)",
    "ECE": "Electronics Engineering",
    "CHE": "Chemical Engineering",
    "EE": "Electrical Engineering",
    "SOD": "School of Design",
    "Sheet13": "Environment & Sustainability Engineering",
    "Physics": "Physics",
}

# data/ss/current_courses.png (code, name, L-T-P, credits, slot, category)
CURRENT_COURSES = [
    ("CSL3090", "Artificial Intelligence", "3-0-0", 3, "B", "PE", "Computer Science & Engineering"),
    ("CSD4030", "B Tech Project", "0-0-6", 3, "T1", "PP", "Computer Science & Engineering"),
    ("LAL3020", "Popular classics", "3-0-0", 3, "J", "OE", "School of Liberal Arts (Humanities)"),
    ("CSN3020", "Design Practical Experience", "0-0-4", 2, "NG3", "ND", "Computer Science & Engineering"),
    ("CSL7650", "Autonomous Systems", "3-0-0", 3, "E", "PE", "Computer Science & Engineering"),
    ("CIL4010", "Environmental Science", "2-0-0", 2, "V", "IS", "Civil and Infrastructure Engineering"),
    ("AIL6110", "Decision Sciences Applications", "3-0-0", 3, "A", "IH", "School of Artificial Intelligence & Data Science"),
    ("MSL4010", "Introduction to Business and Management", "3-0-0", 3, "U", "OE", "School of Management & Entrepreneurship"),
]

ELECTIVE_TOKENS = {"PE", "OE", "SE", "ME", "MO", "LE"}
OTHER_TOKENS = {"PC", "IE", "IS", "IH", "LS", "SC", "NE", "NH", "ND", "PP", "MC", "MP", "IC"}


def clean(v):
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v)).strip()


def first_number(text):
    m = re.search(r"\d+(?:\.\d+)?", text or "")
    return m.group(0) if m else ""


def credits_from(credit_cell, ltp_cell):
    c = clean(credit_cell)
    if c:
        paren = re.search(r"\((\d+(?:\.\d+)?)\)", c)
        if paren:
            return paren.group(1)
        if re.fullmatch(r"\d+(?:\.\d+)?", c):
            return str(int(float(c))) if float(c).is_integer() else c
        if re.match(r"\d+(?:\.\d+)?\s*\(\d", c):  # "3 (3-0-0)"
            return first_number(c)
    l = clean(ltp_cell)
    paren = re.search(r"\((\d+(?:\.\d+)?)\)", l)
    return paren.group(1) if paren else ""


def level_from(code):
    m = re.match(r"[A-Za-z]{2,4}(\d)", code or "")
    return f"{m.group(1)}00" if m else ""


def normalise_type(raw):
    text = clean(raw)
    if not text:
        return ""
    if re.search(r"elective", text, re.I):
        return "Elective"
    tokens = set(re.findall(r"\b[A-Z]{2}\b", text))
    if tokens & ELECTIVE_TOKENS:
        return "Elective"
    if re.search(r"core|compulsory", text, re.I) or tokens & OTHER_TOKENS:
        return "Core"
    return ""


def is_header(cells):
    low = [clean(c).lower() for c in cells]
    has_code = any("code" in c for c in low)
    has_name = any(re.search(r"name|title|subject|^course$", c) for c in low)
    return has_code and has_name


def map_columns(cells):
    low = [clean(c).lower() for c in cells]
    m = {}

    def find(pred, key):
        for i, c in enumerate(low):
            if key not in m and c and pred(c):
                m[key] = i

    find(lambda c: "code" in c, "code")
    find(lambda c: re.search(r"name|title|subject", c), "name")
    if "name" not in m:
        find(lambda c: c == "course", "name")
    find(lambda c: "l-t-p" in c or c == "ltp", "ltp")
    find(lambda c: "credit" in c, "credits")
    find(lambda c: re.search(r"instructor|faculty|coordinator", c) and "co-" not in c, "instructor")
    find(lambda c: c in ("slot", "slots") or ("slot" in c and "lab" not in c), "slot")
    find(lambda c: re.search(r"course type|^type$|subject category", c), "type")
    find(lambda c: re.search(r"offered for|remarks|program|area", c), "notes")
    find(lambda c: "floated" in c, "floated")
    if "slot" not in m and "instructor" in m:
        nxt = m["instructor"] + 1
        if nxt < len(cells) and clean(cells[nxt]) == "" :
            m["slot"] = nxt  # Chemistry sheet leaves the slot header blank
    return m


def parse_sheet(ws, department):
    rows = []
    cols = None
    section = ""
    for cells in ws.iter_rows(values_only=True):
        if not any(clean(c) for c in cells):
            continue
        if is_header(cells):
            cols = map_columns(cells)
            section = ""
            continue
        if cols is None:
            continue

        def get(key):
            i = cols.get(key)
            return clean(cells[i]) if i is not None and i < len(cells) else ""

        name = get("name")
        code = get("code")
        filled = [clean(c) for c in cells if clean(c)]
        if len(filled) <= 2 and not code:  # a section heading such as "Core Courses" or "HSS-2"
            section = " ".join(filled)
            continue
        if not name or name.lower() in ("course name", "course title"):
            continue
        if cols.get("floated") is not None and get("floated").upper() != "Y":
            continue
        code = code.split(",")[0].strip() if re.match(r"^[A-Z]{3}\d", code.split(",")[0].strip()) else code
        if not re.match(r"^[A-Za-z]{2,4}\d", code):
            code = ""  # sheet placeholders such as "(blank)" or serial numbers
        raw_type = get("type")
        notes = get("notes")
        ctype = normalise_type(raw_type) or normalise_type(notes)
        rows.append({
            "course_code": code,
            "course_name": name,
            "department": department,
            "level": level_from(code),
            "credits": credits_from(cells[cols["credits"]] if "credits" in cols else "", get("ltp")),
            "type": ctype,
            "category": raw_type or section,
            "offered_for": notes,
            "semester": SEMESTER,
            "instructor_name": get("instructor"),
            "slot": get("slot"),
        })
    return rows


def load_instructors():
    with open(CSV_DIR / "instructors.csv", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def name_tokens(s):
    s = re.sub(r"\b(dr|prof|professor|mr|ms|mrs)\b\.?", " ", s, flags=re.I)
    return [t for t in re.findall(r"[a-z]+", s.lower()) if len(t) > 1]


def match_instructors(raw, faculty):
    """Return the faculty ids named in a free-text instructor cell (comma, &, /, + separated)."""
    ids = []
    for part in re.split(r"[,&/+;]| and ", raw):
        toks = set(name_tokens(part))
        if not toks:
            continue
        exact = [f["instructor_id"] for f in faculty if set(name_tokens(f["name"])) == toks]
        best = list(exact)
        if not best:  # fall back to partial matches only when no exact name exists
            for f in faculty:
                ftoks = set(name_tokens(f["name"]))
                if (len(toks) >= 2 and (toks <= ftoks or ftoks <= toks)) or (len(toks) == 1 and toks <= ftoks):
                    best.append(f["instructor_id"])
        if len(best) == 1 and best[0] not in ids:
            ids.append(best[0])
    return ids


def slot_times(slot_rows, slot):
    if not re.fullmatch(r"[A-KNUV]", slot or ""):
        return ""
    return "; ".join(f"{r['day']} {r['time']}" for r in slot_rows if r["slot"] == slot)


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    faculty = load_instructors()
    with open(CSV_DIR / "slot_system.csv", encoding="utf-8") as f:
        slot_rows = list(csv.DictReader(f))

    courses = []
    for ws in wb:
        dept = SHEET_DEPT.get(ws.title)
        if dept:
            courses += parse_sheet(ws, dept)
    seen = {(c["course_code"], c["course_name"].lower(), c["department"]) for c in courses}
    for code, name, ltp, credits, slot, category, dept in CURRENT_COURSES:
        if (code, name.lower(), dept) in seen:
            continue
        seen.add((code, name.lower(), dept))
        courses.append({
            "course_code": code, "course_name": name, "department": dept, "level": level_from(code),
            "credits": str(credits), "type": normalise_type(category), "category": category,
            "offered_for": "current_courses.png", "semester": SEMESTER, "instructor_name": "", "slot": slot,
        })

    out_cols = ["course_code", "course_name", "department", "level", "credits", "type", "category",
                "offered_for", "semester", "instructor_name", "instructor_id", "slot"]
    for c in courses:
        c["instructor_id"] = ";".join(match_instructors(c["instructor_name"], faculty))
    with open(CSV_DIR / "courses.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=out_cols)
        w.writeheader()
        w.writerows(courses)

    with open(CSV_DIR / "timetable.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["course_code", "course_name", "department", "slot", "days_time"])
        for c in courses:
            if c["slot"]:
                w.writerow([c["course_code"], c["course_name"], c["department"], c["slot"],
                            slot_times(slot_rows, c["slot"])])

    linked = sum(1 for c in courses if c["instructor_id"])
    print(f"courses: {len(courses)}  with slot: {sum(1 for c in courses if c['slot'])}  "
          f"with instructor id: {linked}  with code: {sum(1 for c in courses if c['course_code'])}")
    by = {}
    for c in courses:
        by.setdefault(c["department"], []).append(c)
    for d, rows in sorted(by.items()):
        lv = {}
        for r in rows:
            lv[r["level"] or "?"] = lv.get(r["level"] or "?", 0) + 1
        print(f"  {d}: {len(rows)}  levels={dict(sorted(lv.items()))}")


if __name__ == "__main__":
    main()
